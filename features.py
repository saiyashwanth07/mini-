# features.py — Final Version (for PDF, DOCX, PPTX, TXT)

import os
import re
from PyPDF2 import PdfReader
from docx import Document
from pptx import Presentation

def extract_features(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    features = {}

    # ---------- PDF ----------
    if ext == ".pdf":
        pdf = PdfReader(file_path)
        features["file_size"] = os.path.getsize(file_path)
        features["pages"] = len(pdf.pages)
        features["metadata size"] = len(pdf.metadata or {})
        features["xref Length"] = len(getattr(pdf, "xref", []))
        features["title characters"] = len(getattr(pdf.metadata, "title", "") or "")
        features["isEncrypted"] = int(pdf.is_encrypted)
        features.update({
            "embedded files": 0, "images": 0, "JS": 0, "Javascript": 0,
            "OpenAction": 0, "Acroform": 0,
            "url_count": 0, "macro_keyword_count": 0, "suspicious_api_count": 0
        })

    # ---------- DOCX ----------
    elif ext == ".docx":
        doc = Document(file_path)
        features["file_size"] = os.path.getsize(file_path)
        features["pages"] = len(doc.paragraphs)
        features.update({
            "metadata size": 0, "xref Length": 0, "title characters": 0, "isEncrypted": 0,
            "embedded files": 0, "images": 0, "JS": 0, "Javascript": 0,
            "OpenAction": 0, "Acroform": 0,
            "url_count": 0, "macro_keyword_count": 0, "suspicious_api_count": 0
        })

    # ---------- PPTX ----------
    elif ext == ".pptx":
        ppt = Presentation(file_path)
        features["file_size"] = os.path.getsize(file_path)
        features["pages"] = len(ppt.slides)
        features.update({
            "metadata size": 0, "xref Length": 0, "title characters": 0, "isEncrypted": 0,
            "embedded files": 0, "images": 0, "JS": 0, "Javascript": 0,
            "OpenAction": 0, "Acroform": 0,
            "url_count": 0, "macro_keyword_count": 0, "suspicious_api_count": 0
        })

    # ---------- TXT ----------
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().lower()

        url_count = len(re.findall(r"http[s]?://", text))
        macro_keywords = ["macro", "vba", "autoopen", "script", "powershell", "cmd.exe", "onload"]
        macro_keyword_count = sum(k in text for k in macro_keywords)
        api_keywords = ["winexec", "shellexecute", "loadlibrary", "urlmon", "createprocess"]
        suspicious_api_count = sum(k in text for k in api_keywords)

        features = {
            "file_size": os.path.getsize(file_path),
            "metadata size": 0,
            "pages": text.count("\n") + 1,
            "xref Length": 0,
            "title characters": 0,
            "isEncrypted": 0,
            "embedded files": 0,
            "images": 0,
            "JS": 0,
            "Javascript": 0,
            "OpenAction": 0,
            "Acroform": 0,
            "url_count": url_count,
            "macro_keyword_count": macro_keyword_count,
            "suspicious_api_count": suspicious_api_count,
        }

    else:
        raise ValueError("Unsupported file type.")

    return list(features.values()), None
